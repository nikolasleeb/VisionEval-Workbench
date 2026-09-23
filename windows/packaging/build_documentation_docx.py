#!/usr/bin/env python3
"""Build editable Windows 2.0 User Guide and What's New DOCX sources."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "user"
GUIDE_OUTPUT = SOURCE / "VisionEval-Workbench-2.0-Windows-User-Guide.docx"
WHATS_NEW_OUTPUT = SOURCE / "VisionEval-Workbench-2.0-Whats-New-Windows.docx"
GUIDE_PAGES = [
    "README.md", "setup.md", "getting-started.md", "core-concepts.md", "explore.md",
    "create-and-review.md", "run.md", "compare.md", "hypercube.md",
    "settings-workspaces-storage.md", "data-units-provenance.md", "troubleshooting.md",
    "keyboard-shortcuts.md", "future-improvements.md",
]
BLUE = RGBColor(23, 54, 93)
MID_BLUE = RGBColor(36, 91, 136)
MUTED = RGBColor(82, 97, 115)


def shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    properties.append(fill)


def page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure(document: Document, title: str) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.72), Inches(0.68)
    section.left_margin = section.right_margin = Inches(0.76)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Aptos", Pt(9.5), RGBColor(31, 41, 51)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(6), 1.12
    for name, size, color in (("Title", 27, BLUE), ("Heading 1", 18, BLUE), ("Heading 2", 14, MID_BLUE), ("Heading 3", 11.5, MID_BLUE)):
        style = styles[name]
        style.font.name, style.font.size, style.font.bold, style.font.color.rgb = "Aptos Display", Pt(size), True, color
        style.paragraph_format.keep_with_next = True
    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name, code.font.size = "Cascadia Mono", Pt(8)
        code.paragraph_format.left_indent = Inches(0.14)
        code.paragraph_format.right_indent = Inches(0.14)
        code.paragraph_format.space_after = Pt(7)
    header = section.header.paragraphs[0]
    header.text = title
    header.style = styles["Caption"]
    header.runs[0].font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("VisionEval Workbench 2.0 — Windows x64   •   ")
    page_field(footer)
    for run in footer.runs:
        run.font.name, run.font.size, run.font.color.rgb = "Aptos", Pt(8), MUTED
    document.core_properties.title = title
    document.core_properties.author = "VisionEval Workbench"


def add_inline(paragraph, text: str) -> None:
    text = re.sub(r"!\[([^]]*)\]\([^)]+\)", "", text)
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name, run.font.size = "Cascadia Mono", Pt(8.5)
        else:
            paragraph.add_run(token)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=width)
    table.style = "Light Shading Accent 1"
    for index, value in enumerate(rows[0]):
        add_inline(table.rows[0].cells[index].paragraphs[0], value)
        shade(table.rows[0].cells[index], "17365D")
        for run in table.rows[0].cells[index].paragraphs[0].runs:
            run.font.bold, run.font.color.rgb = True, RGBColor(255, 255, 255)
    for values in rows[2:]:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            add_inline(cells[index].paragraphs[0], value)


def add_markdown(document: Document, text: str, *, drop_first_title: bool = False) -> None:
    lines = text.splitlines()
    index, first_heading, in_code = 0, True, False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index].rstrip()
        if line.strip().startswith("```"):
            if in_code:
                document.add_paragraph("\n".join(code_lines), "Code Block")
                code_lines.clear()
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[index + 1]):
            rows = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append([item.strip() for item in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            if first_heading and drop_first_title:
                first_heading = False
                index += 1
                continue
            first_heading = False
            paragraph = document.add_paragraph(style=f"Heading {len(heading.group(1))}")
            add_inline(paragraph, heading.group(2))
        elif re.match(r"^[-*]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^[-*]\s+", "", line))
        elif numbered := re.match(r"^(\d+)\.\s+(.+)$", line):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.24)
            paragraph.paragraph_format.first_line_indent = Inches(-0.24)
            add_inline(paragraph, f"{numbered.group(1)}. {numbered.group(2)}")
        elif line.strip() and not line.lstrip().startswith("![") and line.strip() != "---":
            paragraph = document.add_paragraph()
            add_inline(paragraph, line.strip())
        index += 1


def cover(document: Document, subtitle: str) -> None:
    document.add_paragraph("VisionEval Workbench 2.0", "Title")
    paragraph = document.add_paragraph(subtitle)
    paragraph.style = document.styles["Subtitle"]
    paragraph.runs[0].font.color.rgb = MID_BLUE
    status = document.add_paragraph("UNSIGNED RELEASE CANDIDATE")
    status.runs[0].bold, status.runs[0].font.color.rgb = True, RGBColor(155, 58, 34)
    document.add_paragraph("Native Windows x64 • R 4.5.3 • VisionEval VE-40-RC7")


def build_guide() -> None:
    document = Document()
    configure(document, "VisionEval Workbench 2.0 Windows User Guide")
    cover(document, "Windows x64 User Guide")
    document.add_paragraph("This editable source accompanies the verified PDF included with the release candidate.")
    document.add_page_break()
    document.add_heading("Contents", level=1)
    for filename in GUIDE_PAGES:
        source = (SOURCE / filename).read_text(encoding="utf-8")
        heading = next((line[2:] for line in source.splitlines() if line.startswith("# ")), filename)
        document.add_paragraph(heading, "List Bullet")
    for page_index, filename in enumerate(GUIDE_PAGES):
        if page_index:
            document.add_section(WD_SECTION.NEW_PAGE)
        add_markdown(document, (SOURCE / filename).read_text(encoding="utf-8"))
    document.save(GUIDE_OUTPUT)


def build_whats_new() -> None:
    document = Document()
    configure(document, "What's New in VisionEval Workbench 2.0 for Windows")
    cover(document, "What's New for Windows")
    add_markdown(document, (SOURCE / "whats-new.md").read_text(encoding="utf-8"), drop_first_title=True)
    document.save(WHATS_NEW_OUTPUT)


if __name__ == "__main__":
    build_guide()
    build_whats_new()
    print(GUIDE_OUTPUT)
    print(WHATS_NEW_OUTPUT)
